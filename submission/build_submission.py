from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("Security_Portfolio_Submission.docx")
NAVY = "123E36"
INK = "10201D"
MUTED = "58706A"
MOSS = "E1EDD6"
CORAL = "F16B50"
PALE = "F2F4F7"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="C9D4D0", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def fix_table(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_run(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_label_value(doc, label, value, url=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    left = p.add_run(label + "  ")
    set_run(left, size=10.5, color=MUTED, bold=True)
    if url:
        add_hyperlink(p, value, url)
    else:
        r = p.add_run(value)
        set_run(r, size=10.5, color=INK)
    return p


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.492)
    section.footer_distance = Inches(.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (("Heading 1", 16, "2E74B5", 16, 8), ("Heading 2", 13, "2E74B5", 12, 6), ("Heading 3", 12, "1F4D78", 8, 4)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("SECURITY PORTFOLIO SUBMISSION")
    set_run(r, size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Your Name  |  Cybersecurity Portfolio")
    set_run(r, size=8.5, color=MUTED)


def add_checklist_item(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.5)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run(r, size=10.5)


def build():
    doc = Document()
    style_document(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("CAREER PORTFOLIO")
    set_run(r, size=9, color=CORAL, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Cybersecurity Portfolio\nSubmission")
    set_run(r, size=27, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("A concise record of my public security portfolio and LinkedIn placement.")
    set_run(r, size=12, color=MUTED)

    meta = doc.add_table(rows=3, cols=2)
    fix_table(meta, [2700, 6660])
    values = [("Candidate", "Your Name"), ("Date", "August 2026"), ("Portfolio focus", "Threat detection, vulnerability management, network defense, and incident response")]
    for row, (label, value) in zip(meta.rows, values):
        set_cell_shading(row.cells[0], PALE)
        for cell in row.cells:
            set_cell_border(cell)
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run(p1.add_run(label), size=10, color=NAVY, bold=True)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(value), size=10)

    doc.add_heading("Public links", level=1)
    add_label_value(doc, "Live site", "https://YOUR-USERNAME.github.io/security-portfolio/", "https://YOUR-USERNAME.github.io/security-portfolio/")
    add_label_value(doc, "GitHub repository", "https://github.com/YOUR-USERNAME/security-portfolio", "https://github.com/YOUR-USERNAME/security-portfolio")
    add_label_value(doc, "LinkedIn profile", "https://www.linkedin.com/in/YOUR-LINKEDIN-HANDLE/", "https://www.linkedin.com/in/YOUR-LINKEDIN-HANDLE/")

    doc.add_heading("What this portfolio demonstrates", level=1)
    doc.add_paragraph("The linked homepage presents a recruiter-friendly summary of my security focus areas and a body of evidence that grows as I complete labs and projects. Each project is structured around an objective, method, tooling, outcome, and key learning.")
    add_checklist_item(doc, "Hands-on work in security monitoring and investigation")
    add_checklist_item(doc, "Risk-based vulnerability review and remediation planning")
    add_checklist_item(doc, "Network reconnaissance, traffic analysis, and defense fundamentals")

    doc.add_heading("LinkedIn proof", level=1)
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(7)
    r = note.add_run("Before submitting: ")
    set_run(r, size=10.5, color=INK, bold=True)
    r = note.add_run("replace the three placeholders above with live URLs and insert a screenshot below showing the live site added in your LinkedIn Featured section.")
    set_run(r, size=10.5, color=INK)

    proof = doc.add_table(rows=1, cols=1)
    fix_table(proof, [9360])
    cell = proof.cell(0, 0)
    set_cell_shading(cell, "FBFCF9")
    set_cell_border(cell, color="91AFA6", size="12")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("INSERT LINKEDIN FEATURED SCREENSHOT HERE")
    set_run(r, size=13, color=NAVY, bold=True)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(65)
    r = p.add_run("Show the live portfolio card or URL on your profile.")
    set_run(r, size=10.5, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Final check: ")
    set_run(r, size=9.5, color=MUTED, bold=True)
    r = p.add_run("open each link from the exported document, confirm the site is public, and make sure the screenshot includes your name and/or profile context.")
    set_run(r, size=9.5, color=MUTED)

    doc.core_properties.author = "Your Name"
    doc.core_properties.title = "Cybersecurity Portfolio Submission"
    doc.core_properties.subject = "Public portfolio and LinkedIn proof"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
